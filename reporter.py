"""
reporter.py
Write a multi-sheet formatted Excel workbook:
  Summary    — meta stats + grouped totals + bar chart
  Data       — clean full dataset
  Quality    — per-column missing-rate table
  Exceptions — rows with at least one anomaly flag (incl. conversion failures)
"""
from __future__ import annotations

import os

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_BLUE_FILL = PatternFill("solid", fgColor="2E5A87")
_BLUE_FONT = Font(color="FFFFFF", bold=True)
_WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
_WARN_FONT = Font(color="9C5700", bold=True)


def _style_header(ws, ncols: int, row: int = 1, fill=_BLUE_FILL, font=_BLUE_FONT) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max(
            (len(str(c.value)) for c in col if c.value is not None), default=8
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 40)


def build_summaries(
    df: pd.DataFrame, amount_col: str, group_cols: list[str]
) -> dict[str, pd.DataFrame]:
    """Build one summary table per group column found in df."""
    summaries: dict[str, pd.DataFrame] = {}
    for group in group_cols:
        if group not in df.columns or amount_col not in df.columns:
            continue
        tbl = (
            df.groupby(group)[amount_col]
            .agg(["sum", "count"])
            .rename(columns={"sum": f"Total {amount_col}", "count": "Rows"})
            .sort_values(f"Total {amount_col}", ascending=False)
            .reset_index()
        )
        tbl[f"Total {amount_col}"] = tbl[f"Total {amount_col}"].round(2)
        summaries[group] = tbl
    return summaries


def write_report(
    df: pd.DataFrame,
    quality: pd.DataFrame,
    summaries: dict[str, pd.DataFrame],
    conversion_failures: dict[str, int],
    failed_files: list[str],
    amount_col: str,
    out_path: str,
    n_files: int,
) -> None:
    """Write the four-sheet workbook to out_path."""
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)

    # Exceptions = any row with a non-empty _anomaly_flags
    flag_col = df.get("_anomaly_flags", pd.Series("", index=df.index))
    exceptions = df[flag_col != ""].copy()

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:

        # ── Summary ───────────────────────────────────────────────────────────
        total_amt = (
            round(df[amount_col].sum(), 2) if amount_col in df.columns else "n/a"
        )
        conv_fail_str = (
            "; ".join(f"{c}: {n} row(s)" for c, n in conversion_failures.items())
            if conversion_failures else "none"
        )
        failed_str = (
            "; ".join(failed_files) if failed_files else "none"
        )
        meta = pd.DataFrame({
            "Metric": [
                "Files read",
                "Files failed (skipped)",
                "Clean rows",
                f"Total {amount_col}",
                "Anomaly / flagged rows",
                "Type conversion failures",
            ],
            "Value": [
                n_files,
                failed_str,
                len(df),
                total_amt,
                len(exceptions),
                conv_fail_str,
            ],
        })
        start = 0
        meta.to_excel(writer, sheet_name="Summary", startrow=start, index=False)
        start += len(meta) + 3

        chart_anchor: tuple[int, int] | None = None
        for name, tbl in summaries.items():
            tbl.to_excel(writer, sheet_name="Summary", startrow=start, index=False)
            if chart_anchor is None and amount_col in df.columns:
                chart_anchor = (start, len(tbl))
            start += len(tbl) + 3

        # ── Data ──────────────────────────────────────────────────────────────
        df.to_excel(writer, sheet_name="Data", index=False)

        # ── Quality ───────────────────────────────────────────────────────────
        quality.to_excel(writer, sheet_name="Quality", index=False)

        # ── Exceptions ────────────────────────────────────────────────────────
        if not exceptions.empty:
            exceptions.to_excel(writer, sheet_name="Exceptions", index=False)
        else:
            pd.DataFrame({"Note": ["No anomalies or conversion failures detected."]}).to_excel(
                writer, sheet_name="Exceptions", index=False
            )

        # ── Formatting ────────────────────────────────────────────────────────
        wb = writer.book

        ws_sum = wb["Summary"]
        _style_header(ws_sum, meta.shape[1], row=1)
        row_cursor = len(meta) + 3 + 1
        for tbl in summaries.values():
            _style_header(ws_sum, tbl.shape[1], row=row_cursor)
            row_cursor += len(tbl) + 3
        _autosize(ws_sum)

        ws_data = wb["Data"]
        _style_header(ws_data, df.shape[1], row=1)
        ws_data.freeze_panes = "A2"
        _autosize(ws_data)

        ws_q = wb["Quality"]
        _style_header(ws_q, quality.shape[1], row=1, fill=_WARN_FILL, font=_WARN_FONT)
        _autosize(ws_q)

        ws_ex = wb["Exceptions"]
        _style_header(ws_ex, ws_ex.max_column, row=1, fill=_WARN_FILL, font=_WARN_FONT)
        _autosize(ws_ex)

        # ── Bar chart ─────────────────────────────────────────────────────────
        if chart_anchor and summaries:
            header_row0, nrows = chart_anchor
            first_data = header_row0 + 2
            last_data  = header_row0 + 1 + nrows
            chart = BarChart()
            chart.title = f"Total {amount_col} by {next(iter(summaries))}"
            chart.y_axis.title = amount_col
            chart.x_axis.title = next(iter(summaries))
            data_ref = Reference(ws_sum, min_col=2, min_row=header_row0 + 1, max_row=last_data)
            cats_ref = Reference(ws_sum, min_col=1, min_row=first_data,      max_row=last_data)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.height = 8
            chart.width  = 16
            ws_sum.add_chart(chart, "E1")
